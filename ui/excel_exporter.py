"""监控结果 Excel 导出（**仅依赖 openpyxl**）。

为什么不用 pandas：本项目只做「追加行 + 自适应列宽」这一件事，pandas 会连带
numpy 一起进打包（约 40~60MB），对离线分发的体积影响远大于收益。openpyxl 已
是 requirements 的依赖，直接读写即可，行为与旧实现一致。
"""

import os
from datetime import datetime

from config import EXCEL_DIR
from service.logging_service import get_logger

MAX_COLUMN_WIDTH = 50
"""列宽上限（与旧实现一致）"""


class ExcelExporter:
    """把监控结果追加写入单个 Excel 文件。"""

    def __init__(self, log_callback=None):
        self._log = log_callback or (lambda *args, **kwargs: None)
        self._logger = get_logger(__name__)
        self.excel_file = None

    def reset(self) -> None:
        """清除当前文件句柄，下次导出重新生成文件名。"""
        self.excel_file = None

    def export(self, device_data, dev_type, devices, options, format_time) -> str:
        """导出一次监控快照，返回 Excel 路径；失败返回空字符串。

        options 需包含 monitor_time / monitor_boot / monitor_start_count 三个布尔值；
        format_time 用于把启动时间格式化为分钟精度。
        """
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self._log("导出Excel失败：未找到openpyxl模块。请运行 'pip install openpyxl' 安装依赖",
                      is_error=True)
            return ""

        try:
            os.makedirs(EXCEL_DIR, exist_ok=True)
            if not self.excel_file:
                timestamp = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
                self.excel_file = os.path.join(EXCEL_DIR, f'{dev_type}_{timestamp}.xlsx')

            monitor_time = bool(options.get('monitor_time'))
            monitor_boot = bool(options.get('monitor_boot'))
            monitor_start_count = bool(options.get('monitor_start_count'))

            current_system_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            device_index = {d.get('id'): d for d in devices if isinstance(d, dict)}

            rows = []
            for device_id, info in device_data.items():
                device_info = device_index.get(device_id, {})
                row = {
                    '系统时间': current_system_time,
                    '设备ID': device_id,
                    'IP地址': device_info.get('ip', '未知'),
                    'MAC地址': device_info.get('mac_address', '未知'),
                    '设备名称': device_info.get('device_name', '未知'),
                }
                if monitor_time:
                    row['设备时间'] = info.get('last_card_time', info.get('initial_card_time', 'N/A'))
                    time_diff = info.get('last_time_diff', info.get('initial_time_diff', 'N/A'))
                    row['时间差'] = f"{time_diff}秒" if time_diff != 'N/A' else time_diff
                if monitor_boot:
                    row['启动时间'] = format_time(info.get('last_boot_time', 'N/A'))
                    row['运行时长'] = info.get('last_uptime', 'N/A')
                if monitor_start_count:
                    row['播放器启动次数'] = info.get('last_start_count', 'N/A')
                rows.append(row)

            if not rows:
                return self.excel_file

            file_exists = os.path.exists(self.excel_file)
            workbook = load_workbook(self.excel_file) if file_exists else Workbook()
            worksheet = workbook.active
            if not file_exists:
                worksheet.title = '设备监控数据'

            self._append_rows(worksheet, rows)
            workbook.save(self.excel_file)
            workbook.close()
        except Exception as e:
            self._logger.exception("写入Excel文件时出错")
            self._log(f"写入Excel文件时出错: {str(e)}", is_error=True)
            return ""

        excel_path = os.path.abspath(self.excel_file).replace('\\', '/')
        self._log(f"监控结果已成功导出到 {excel_path}")
        return self.excel_file

    @staticmethod
    def _append_rows(worksheet, rows) -> None:
        """按表头对齐追加行；出现新列时自动扩列并同步列宽。"""
        headers = []
        if worksheet.max_row >= 1:
            headers = [cell.value for cell in worksheet[1]]
            while headers and headers[-1] is None:   # 去掉表头尾部的空单元格
                headers.pop()

        if not headers:
            # 新文件：用首行的键作为表头（保持插入顺序）
            headers = list(rows[0].keys())
            for column, name in enumerate(headers, start=1):
                worksheet.cell(row=1, column=column, value=name)
                worksheet.column_dimensions[
                    worksheet.cell(row=1, column=column).column_letter].width = max(
                        12, min(len(str(name)) + 15, MAX_COLUMN_WIDTH))

        for raw in rows:
            for name in raw:
                if name not in headers:
                    headers.append(name)
                    column = len(headers)
                    worksheet.cell(row=1, column=column, value=name)
            values = [raw.get(name, '') for name in headers]
            worksheet.append(values)

            # 只更新本轮新增值更长的列（不重算全表，避免大文件反复遍历）
            for column, value in enumerate(values, start=1):
                text = '' if value is None else str(value)
                letter = worksheet.cell(row=1, column=column).column_letter
                current = worksheet.column_dimensions[letter].width or 0
                wanted = min(max(len(text) + 4, 12), MAX_COLUMN_WIDTH)
                if wanted > current:
                    worksheet.column_dimensions[letter].width = wanted
