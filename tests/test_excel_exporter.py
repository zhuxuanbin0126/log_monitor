"""Excel 导出（openpyxl 版）：追加语义、列宽自适应、表头扩展。

不依赖 pandas/numpy —— 这两者只为「追加行」进包，体积代价远大于收益。
"""

import os

from openpyxl import load_workbook

from ui.excel_exporter import ExcelExporter


def _format_time(value):
    return value or 'N/A'


def _device_data():
    return {
        'H16-26-A0419': {
            'last_card_time': '2026-10-20 15:00:00',
            'last_time_diff': 1,
            'last_boot_time': '2026-10-20T08:00:00',
            'last_uptime': '7时0分',
            'last_start_count': '1',
        }
    }


def test_creates_file_with_header_and_row(monkeypatch, tmp_path):
    monkeypatch.setattr("ui.excel_exporter.EXCEL_DIR", str(tmp_path))
    exporter = ExcelExporter()
    path = exporter.export(_device_data(), 'H16', [{'id': 'H16-26-A0419', 'ip': '1.1.1.1'}],
                           {'monitor_time': True, 'monitor_boot': True,
                            'monitor_start_count': True}, _format_time)

    assert path and os.path.exists(path)
    sheet = load_workbook(path).active
    headers = [cell.value for cell in sheet[1]]
    assert headers[:3] == ['系统时间', '设备ID', 'IP地址']
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=2).value == 'H16-26-A0419'


def test_appends_without_duplicating_header(monkeypatch, tmp_path):
    monkeypatch.setattr("ui.excel_exporter.EXCEL_DIR", str(tmp_path))
    exporter = ExcelExporter()
    options = {'monitor_time': True, 'monitor_boot': False, 'monitor_start_count': False}
    exporter.export(_device_data(), 'H16', [], options, _format_time)
    exporter.export(_device_data(), 'H16', [], options, _format_time)

    sheet = load_workbook(exporter.excel_file).active
    assert sheet.max_row == 3                     # 表头 + 2 行数据
    assert len([c for c in sheet[1] if c.value]) == 7


def test_new_column_extends_existing_header(monkeypatch, tmp_path):
    monkeypatch.setattr("ui.excel_exporter.EXCEL_DIR", str(tmp_path))
    exporter = ExcelExporter()
    exporter.export(_device_data(), 'H16', [],
                    {'monitor_time': False, 'monitor_boot': False,
                     'monitor_start_count': False}, _format_time)
    exporter.export(_device_data(), 'H16', [],
                    {'monitor_time': False, 'monitor_boot': False,
                     'monitor_start_count': True}, _format_time)

    sheet = load_workbook(exporter.excel_file).active
    headers = [cell.value for cell in sheet[1]]
    assert '播放器启动次数' in headers
    # 老行该列留空，新行有值
    column = headers.index('播放器启动次数') + 1
    assert sheet.cell(row=2, column=column).value in (None, '')
    assert sheet.cell(row=3, column=column).value == '1'


def test_reset_starts_new_file(monkeypatch, tmp_path):
    monkeypatch.setattr("ui.excel_exporter.EXCEL_DIR", str(tmp_path))
    exporter = ExcelExporter()
    options = {'monitor_time': True, 'monitor_boot': False, 'monitor_start_count': False}
    first = exporter.export(_device_data(), 'H16', [], options, _format_time)
    exporter.reset()
    second = exporter.export(_device_data(), 'H16', [], options, _format_time)

    assert first and second
    assert os.path.exists(first) and os.path.exists(second)


def test_empty_rows_keeps_file_path(monkeypatch, tmp_path):
    monkeypatch.setattr("ui.excel_exporter.EXCEL_DIR", str(tmp_path))
    exporter = ExcelExporter()
    path = exporter.export({}, 'H16', [], {'monitor_time': True}, _format_time)
    assert path.endswith('.xlsx')
